#!/usr/bin/env python3
"""Optional scaffold for Week 1 learning journal (Spanish).

Canonical student file: course/week-1/learning-journal-week-1.docx
(edit in Word). Regenerating overwrites that path — merge manual edits
back into this script if you rely on regen.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Pt

REPO_ROOT = Path(__file__).resolve().parents[2]
OUT_PATH = (
    REPO_ROOT
    / "work-space/teaching/big-data/course/week-1/learning-journal-week-1.docx"
)


def _style_body(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.15
    pf.space_after = Pt(6)


def _heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def _para(doc: Document, text: str, *, bold: bool = False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold


def _question(doc: Document, number: str, text: str) -> None:
    p = doc.add_paragraph()
    p.add_run(f"{number}. ").bold = True
    p.add_run(text)
    doc.add_paragraph("Respuesta:", style="Intense Quote")
    for _ in range(3):
        doc.add_paragraph()


def build() -> Path:
    doc = Document()
    _style_body(doc)

    _heading(doc, "Diario de aprendizaje — Semana 1", 0)
    _para(doc, "Maestría en Estadística Aplicada · UDENAR · A-2026")
    _para(doc, "Electiva I · Big Data")
    _para(doc, "Lecturas previas al sábado 6 de junio de 2026 (L1 + L2)")

    doc.add_paragraph()
    _heading(doc, "Instrucciones", 1)
    for line in [
        "Este cuaderno es personal.",
        "Su función es guiar la lectura de los tres textos antes del sábado. "
        "Lleve sus respuestas a la discusión grupal de lecturas y a la plenaria "
        "(~1 h 15 min).",
        "Escriba en español o en inglés (elija uno y manténgalo en todo el documento).",
        "Extensión orientativa: ~600–900 palabras en total (comprensión + reflexión).",
        "Las lecturas se comparten a través del correo electrónico del curso.",
    ]:
        doc.add_paragraph(line, style="List Bullet")

    _heading(doc, "Lecturas de la semana", 1)
    readings = [
        (
            "boyd & Crawford (2012)",
            "“Critical Questions for Big Data”. Archivo: boyd-crawford-2012.pdf",
        ),
        (
            "Zuboff (2019)",
            "La era del capitalismo de la vigilancia — Capítulo 1. "
            "Archivo: zuboff-es-2019.pdf (edición en español) o "
            "zuboff-en-2019.pdf (edición en inglés).",
        ),
        (
            "Mittelstadt et al. (2016)",
            "“The ethics of algorithms: Mapping the debate”. "
            "Archivo: mittelstadt-et-al-2016.pdf",
        ),
    ]
    for title, detail in readings:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(f"{title}. ").bold = True
        p.add_run(detail)

    _heading(doc, "Parte A — Lectura 1: Introducción a Big Data", 1)
    _para(
        doc,
        "Lea boyd & Crawford y Zuboff Cap. 1. Las preguntas de comprensión "
        "verifican que captó las ideas del texto; las de reflexión temporal "
        "piden comparar el momento en que se escribió con la situación actual "
        "(2026).",
    )

    _heading(doc, "A. boyd & Crawford (2012)", 2)
    _question(
        doc,
        "A1",
        "En dos o tres oraciones: ¿cuál es la tesis principal del artículo? "
        "¿Qué critican los autores del discurso dominante sobre “Big Data”?",
    )
    _question(
        doc,
        "A2",
        "Enumere tres “preguntas críticas” que propone el artículo y explique "
        "una de ellas con sus propias palabras.",
    )
    _question(
        doc,
        "A3",
        "¿Qué distingue el artículo entre “más datos” y “mejor comprensión”? "
        "Cite una idea concreta del texto.",
    )
    _question(
        doc,
        "A4",
        "Responda una o más “preguntas críticas” de acuerdo a la evolución "
        "tecnológica desde la publicación del artículo (i.e., 2012).",
    )
    _question(
        doc,
        "A5",
        "¿Qué limitación o riesgo del “Big Data” señala el artículo que sigue "
        "siendo relevante hoy?",
    )

    _heading(doc, "B. Zuboff — Capítulo 1 (2019)", 2)
    _question(
        doc,
        "B1",
        "En un párrafo: ¿qué es el “capitalismo de la vigilancia” según Zuboff?",
    )
    _question(
        doc,
        "B2",
        "¿Qué relación describe el capítulo entre datos conductuales, predicción "
        "y poder económico? Resuma con sus palabras.",
    )
    _question(
        doc,
        "B3",
        "¿En qué se diferencia, según el texto, mejorar un servicio para el "
        "usuario de extraer valor comercial de los datos generados?",
    )
    _question(
        doc,
        "B4",
        "Anote una afirmación del capítulo que le resulte convincente y otra "
        "que le genere dudas. Explique brevemente cada una.",
    )

    _heading(doc, "Parte B — Lectura 2: Ética y gobernanza de datos", 1)

    _heading(doc, "C. Mittelstadt et al. (2016)", 2)
    _question(
        doc,
        "C1",
        "¿Qué problema central intenta ordenar el artículo? ¿Qué tipo de "
        "“mapa” del debate propone?",
    )
    _question(
        doc,
        "C2",
        "Nombre dos tipos de daño o de preocupación ética del artículo y "
        "explique uno con sus palabras.",
    )
    _question(
        doc,
        "C3",
        "¿Qué actores o partes interesadas menciona el texto (p. ej. "
        "desarrolladores, reguladores, personas afectadas)? ¿Quién falta "
        "en el mapa, si alguien falta?",
    )
    _question(
        doc,
        "C4",
        "Elija un término del artículo (responsabilidad, transparencia, "
        "equidad, autonomía, explicabilidad, etc.) y defínalo como lo "
        "entiende usted después de leer.",
    )

    _heading(doc, "D. Síntesis entre lecturas", 2)
    _question(
        doc,
        "D1",
        "En un párrafo: ¿cómo se complementan boyd & Crawford, Zuboff y "
        "Mittelstadt et al. al pensar en datos a gran escala?",
    )
    _question(
        doc,
        "D2",
        "¿En qué coinciden los tres textos sobre poder, contexto o límites "
        "de los datos? ¿En qué se separan?",
    )
    _question(
        doc,
        "D3",
        "Escriba dos o tres preguntas para la discusión grupal del sábado "
        "sobre las lecturas.",
    )

    _heading(doc, "Notas libres", 1)
    _para(doc, "Espacio opcional durante la discusión del sábado:")
    for _ in range(6):
        doc.add_paragraph()

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_PATH)
    return OUT_PATH


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path}")
