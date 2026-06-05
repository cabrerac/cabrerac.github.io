#!/usr/bin/env python3
"""Build L1 student Word templates (Spanish) for download on the lecture website.

Outputs (canonical for the site):
  assets/documents/26-udenar-big-data/project-requirements-template.docx
  assets/documents/26-udenar-big-data/reflection-week-1-template.docx

Also mirrors to work-space/teaching/big-data/course/week-1/ for instructor editing.

Regenerate after changing section titles (must match check_submission.py).
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Pt

REPO = Path(__file__).resolve().parents[2]
COURSE = "26-udenar-big-data"
ASSETS_DIR = REPO / "assets" / "documents" / COURSE
WEEK1_DIR = REPO / "work-space/teaching/big-data/course/week-1"

REQ_ASSETS = ASSETS_DIR / "project-requirements-template.docx"
REF_ASSETS = ASSETS_DIR / "reflection-week-1-template.docx"
REQ_WEEK1 = WEEK1_DIR / "project-requirements-template.docx"
REF_WEEK1 = WEEK1_DIR / "reflection-week-1-template.docx"


def _style_body(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.15
    pf.space_after = Pt(6)


def _heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def _para(doc: Document, text: str) -> None:
    doc.add_paragraph(text)


def _section(doc: Document, title: str, prompt: str) -> None:
    _heading(doc, title, 2)
    _para(doc, prompt)
    doc.add_paragraph()


def build_project_requirements() -> Document:
    doc = Document()
    _style_body(doc)
    _heading(doc, "Requerimientos del proyecto — Lección 1", 0)
    _para(doc, "Maestría en Estadística Aplicada · UDENAR · Electiva Big Data · A-2026")
    _para(
        doc,
        "Complete en Word, inserte su diagrama de pipeline en la sección 9, "
        "pegue su modelo canvas en la sección 11, "
        "y entregue project_requirements.pdf en Moodle (semana 2; plazo según calendario del curso). "
        "Mantenga los títulos de sección exactamente como aparecen abajo.",
    )
    doc.add_paragraph()

    sections = [
        (
            "1. Metadatos del grupo",
            "Id del grupo, integrantes, arquetipo seleccionado (asignación de recursos | riesgo / alerta temprana | monitoreo).",
        ),
        (
            "2. Interesados",
            "¿A quién sirve su equipo de analítica y quién se ve afectado por las decisiones?",
        ),
        (
            "3. Requerimientos del proyecto",
            "Liste los requerimientos del proyecto que se deben satisfacer para el éxito del proyecto.",
        ),
        (
            "4. Objetivos",
            "Liste los objetivos del proyecto y explique por qué son relevantes.",
        ),
        (
            "5. Subconjunto de datos (GEIH)",
            "Describa las variables del dataset GEIH que usted considera utilizar en su proyecto. Por ejemplo, años (p. ej. 2022–2025), regiones/departamentos, segmento de población, variables principales, etc.",
        ),
        (
            "6. Preocupación ética y lectura",
            "Liste los riesgos o temas éticos concretos (p. ej. divulgación, representación, mezcla GEIH–OSM) que son relevantes para su proyecto y explique por qué son relevantes. "
        ),
        (
            "7. Mecanismos de ética y gobernanza de datos",
            "Liste mecanismos concretos que su pipeline aplicará para mitigar los riesgos "
            "(p. ej. agregación mínima k, no publicar microdatos, documentar cobertura OSM, "
            "manifiesto de acceso, revisión antes de entregables).",
        ),
        (
            "8. Roles y plan de contribución",
            "Una viñeta por integrante (Marco de trabajo, Datos, Análisis, Escritura, Presentación).",
        ),
        (
            "9. Pipeline del proyecto",
            "Inserte una figura que muestre el pipeline del proyecto y 3–5 oraciones que lo describan. Ejemplo: fuentes → almacenamiento → procesamiento → decisión.",
        ),
        (
            "10. Rotación de liderazgo para su grupo por semana",
            "S1: … S2: … S3: … S4: … (cada integrante debe liderar al menos una vez por semana).",
        ),
        (
            "11. Modelo canvas del proyecto",
            "Elabore un diagrama tipo canvas (recuadros con los elementos clave de su proyecto) "
            "que muestre visualmente los requerimientos declarados en las secciones 2–8. "
            "Puede usar la herramienta o formato que prefiera (draw.io, PowerPoint, Miro, Canva, etc.). "
            "Copie y pegue aquí la imagen del canvas completado.",
        ),
    ]
    for title, prompt in sections:
        _section(doc, title, prompt)
    return doc


def build_reflection_week_1() -> Document:
    doc = Document()
    _style_body(doc)
    _heading(doc, "Reflexión individual — Semana 1 (L1 + L2)", 0)
    _para(
        doc,
        "Entregable: week-1-reflection-<student>.pdf (exportar desde este Word). "
        "Plazo: martes 9 de junio de 2026, 23:59 (Colombia). "
        "Cubre Lección 1 (Access) y Lección 2 (ética y gobernanza). Reemplace student por su nombre.",
    )
    _para(doc, "Extensión orientativa del cuerpo: 700–1,000 palabras.")
    doc.add_paragraph()

    _section(
        doc,
        "Metadatos",
        "Semana: 1 (L1 + L2) | Estudiante: … | Grupo: G… | Lecturas citadas: … | Conteo de palabras: ~…",
    )
    _section(
        doc,
        "R1 Proceso y metodología",
        "¿Qué hizo en los cuadernos l1-introduction y l2-ethics-governance? "
        "¿Qué aprendió sobre Access y sobre auditoría ética?",
    )
    _section(
        doc,
        "R2 Justificación técnica",
        "Identifique una decisión técnica (descarga DANE, consulta Overpass u otra) "
        "y explique el por qué.",
    )
    _section(
        doc,
        "R3 Enlace con el dominio y los datos",
        "¿Cómo se relacionan GEIH y OSM con su proyecto grupal?"
    )
    _section(
        doc,
        "R4 Ética y lecturas",
        "Relacione las lecturas con los aspectos éticos y de gobernanza en los ejercicios prácticos de los cuadernos.",
    )
    _section(
        doc,
        "R5 Uso de IA e integridad",
        "¿Qué herramientas de IA utilizó en la semana 1? "
        "¿Cómo influyó en su trabajo y qué aprendió sobre su práctica?",
    )
    _section(
        doc,
        "R6 Retroalimentación sobre la semana",
        "¿Qué mejoraría o qué funcionó bien en esta primera semana del curso? "
        "No se califica; únicamente ayuda a ajustar el curso.",
    )
    return doc


def _save_pair(doc: Document, assets_path: Path, week1_path: Path) -> None:
    assets_path.parent.mkdir(parents=True, exist_ok=True)
    week1_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(assets_path)
    doc.save(week1_path)


def build() -> tuple[Path, Path]:
    req = build_project_requirements()
    ref = build_reflection_week_1()
    _save_pair(req, REQ_ASSETS, REQ_WEEK1)
    _save_pair(ref, REF_ASSETS, REF_WEEK1)
    return REQ_ASSETS, REF_ASSETS


if __name__ == "__main__":
    req_path, ref_path = build()
    print(f"Wrote {req_path}")
    print(f"Wrote {ref_path}")
    print(f"Mirror {REQ_WEEK1}")
    print(f"Mirror {REF_WEEK1}")
