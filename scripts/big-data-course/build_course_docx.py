#!/usr/bin/env python3
"""Build Big Data course Word documents (Spanish).

Regenerate after changing section titles (must match check_submission.py).

Examples:
  python scripts/big-data-course/build_course_docx.py --all
  python scripts/big-data-course/build_course_docx.py reflection_week_2
"""

from __future__ import annotations

import argparse
from collections.abc import Callable
from pathlib import Path

from docx import Document

from docx_build import (
    ASSETS_DIR,
    bullet_list,
    heading,
    para,
    question,
    reading_bullet,
    PROJECT_DIR,
    save_doc,
    save_mirror,
    section,
    style_body,
    WEEK1_DIR,
    WEEK2_DIR,
    WEEK3_DIR,
)

# --- paths -----------------------------------------------------------------

PROJECT_REQUIREMENTS_ASSETS = ASSETS_DIR / "project-requirements-template.docx"
PROJECT_REQUIREMENTS_WEEK1 = WEEK1_DIR / "project-requirements-template.docx"
REFLECTION_WEEK_1_ASSETS = ASSETS_DIR / "reflection-week-1-template.docx"
REFLECTION_WEEK_1_WEEK1 = WEEK1_DIR / "reflection-week-1-template.docx"
REFLECTION_WEEK_2_ASSETS = ASSETS_DIR / "reflection-week-2-template.docx"
REFLECTION_WEEK_2_WEEK2 = WEEK2_DIR / "reflection-week-2-template.docx"
REFLECTION_WEEK_3_ASSETS = ASSETS_DIR / "reflection-week-3-template.docx"
REFLECTION_WEEK_3_WEEK3 = WEEK3_DIR / "reflection-week-3-template.docx"
LEARNING_JOURNAL_WEEK_1 = WEEK1_DIR / "learning-journal-week-1.docx"
LEARNING_JOURNAL_WEEK_1_ASSETS = ASSETS_DIR / "learning-journal-week-1.docx"
LEARNING_JOURNAL_WEEK_2 = WEEK2_DIR / "learning-journal-week-2.docx"
LEARNING_JOURNAL_WEEK_2_ASSETS = ASSETS_DIR / "learning-journal-week-2.docx"
LEARNING_JOURNAL_WEEK_3 = WEEK3_DIR / "learning-journal-week-3.docx"
LEARNING_JOURNAL_WEEK_3_ASSETS = ASSETS_DIR / "learning-journal-week-3.docx"
PROJECT_REPORT_ASSETS = ASSETS_DIR / "project-report-template.docx"
PROJECT_REPORT_PROJECT = PROJECT_DIR / "project-report-template.docx"
PROJECT_REFLECTION_ASSETS = ASSETS_DIR / "project-reflection-template.docx"
PROJECT_REFLECTION_PROJECT = PROJECT_DIR / "project-reflection-template.docx"
DATA_ARCHITECTURE_ASSETS = ASSETS_DIR / "data-architecture-template.docx"
DATA_ARCHITECTURE_WEEK2 = WEEK2_DIR / "data-architecture-template.docx"


# --- document builders -------------------------------------------------------


def build_project_requirements() -> Document:
    doc = Document()
    style_body(doc)
    heading(doc, "Requerimientos del proyecto — Lección 1", 0)
    para(doc, "Maestría en Estadística Aplicada · UDENAR · Electiva Big Data · A-2026")
    para(
        doc,
        "Complete en Word, inserte su diagrama de pipeline en la sección 9, "
        "pegue su modelo canvas en la sección 11. "
        "Avance el documento en la sesión del sábado 13 de junio (checkpoint: subconjunto GEIH definido en §5). "
        "Entregue project_requirements.pdf con el proyecto summativo (sábado 27 de junio de 2026). "
        "Mantenga los títulos de sección exactamente como aparecen abajo.",
    )
    doc.add_paragraph()

    sections_data = [
        (
            "1. Metadatos del grupo",
            "Id del grupo, integrantes, arquetipo seleccionado (asignación de recursos | riesgo / alerta temprana | monitoreo).",
        ),
        (
            "2. Interesados",
            "¿A quién sirve su equipo de analítica y quién se ve afectado por las decisiones?",
        ),
        (
            "3. Requerimientos del proyecto",
            "Liste los requerimientos del proyecto que se deben satisfacer para el éxito del proyecto.",
        ),
        (
            "4. Objetivos",
            "Liste los objetivos del proyecto y explique por qué son relevantes.",
        ),
        (
            "5. Subconjunto de datos (GEIH)",
            "Describa las variables del dataset GEIH que usted considera utilizar en su proyecto. Por ejemplo, años (p. ej. 2022–2025), regiones/departamentos, segmento de población, variables principales, etc.",
        ),
        (
            "6. Preocupación ética y lectura",
            "Liste los riesgos o temas éticos concretos (p. ej. divulgación, representación, mezcla GEIH–OSM) que son relevantes para su proyecto y explique por qué son relevantes. ",
        ),
        (
            "7. Mecanismos de ética y gobernanza de datos",
            "Liste mecanismos concretos que su pipeline aplicará para mitigar los riesgos "
            "(p. ej. agregación mínima k, no publicar microdatos, documentar cobertura OSM, "
            "manifiesto de acceso, revisión antes de entregables).",
        ),
        (
            "8. Roles y plan de contribución",
            "Una viñeta por integrante (Marco de trabajo, Datos, Análisis, Escritura, Presentación).",
        ),
        (
            "9. Pipeline del proyecto",
            "Inserte una figura que muestre el pipeline del proyecto y 3–5 oraciones que lo describan. Ejemplo: fuentes → almacenamiento → procesamiento → decisión.",
        ),
        (
            "10. Rotación de liderazgo para su grupo por semana",
            "S1: … S2: … S3: … S4: … (cada integrante debe liderar al menos una vez por semana).",
        ),
        (
            "11. Modelo canvas del proyecto",
            "Elabore un diagrama tipo canvas (recuadros con los elementos clave de su proyecto) "
            "que muestre visualmente los requerimientos declarados en las secciones 2–8. "
            "Puede usar la herramienta o formato que prefiera (draw.io, PowerPoint, Miro, Canva, etc.). "
            "Copie y pegue aquí la imagen del canvas completado.",
        ),
    ]
    for title, prompt in sections_data:
        section(doc, title, prompt)
    return doc


def build_data_architecture() -> Document:
    doc = Document()
    style_body(doc)
    heading(doc, "Arquitectura de datos del proyecto — Semana 2", 0)
    para(doc, "Maestría en Estadística Aplicada · UDENAR · Electiva Big Data · A-2026")
    para(
        doc,
        "Complete en Word. Pegue su diagrama de arquitectura en la sección 2 y explíquelo en la sección 3. "
        "Use el ejemplo aICU de la Lección 4 como referencia (no copie literalmente). "
        "Avance en la sesión del sábado 13 de junio (bloque de arquitectura de datos). "
        "Entregue data_architecture.pdf con el proyecto summativo (sábado 27 de junio de 2026). "
        "Mantenga los títulos de sección exactamente como aparecen abajo.",
    )
    doc.add_paragraph()

    sections_data = [
        (
            "1. Metadatos del grupo",
            "Id del grupo, integrantes, arquetipo (asignación de recursos | riesgo / alerta temprana | monitoreo).",
        ),
        (
            "2. Diagrama de arquitectura de datos",
            "Inserte aquí una figura con su arquitectura (fuentes → almacenamiento → procesamiento → consumidores). "
            "Herramientas sugeridas: draw.io, PowerPoint, Excalidraw, Miro, Canva, etc.",
        ),
        (
            "3. Explicación del diagrama",
            "Describa cada capa o caja del diagrama: qué datos entran, dónde se almacenan (formato, partición), "
            "qué procesamiento aplica su grupo (harmonización, consultas, agregados) y quién usa los resultados "
            "(tareas de assess y address). 400–700 palabras orientativas.",
        ),
        (
            "4. Relación con los requerimientos del proyecto",
            "Indique cómo esta arquitectura implementa el subconjunto GEIH y el pipeline declarados en "
            "project_requirements (secciones 5 y 9).",
        ),
        (
            "5. Gobernanza y acceso en la arquitectura",
            "¿Quién accede a cada capa? ¿Qué no se publica? ¿Qué mecanismos de ética (§6–§7 de requerimientos) "
            "quedan reflejados en el diseño?",
        ),
    ]
    for title, prompt in sections_data:
        section(doc, title, prompt)
        if title.startswith("2."):
            para(doc, "[Espacio para pegar el diagrama — inserte imagen en Word]")
            for _ in range(4):
                doc.add_paragraph()
    return doc


def build_reflection_week_1() -> Document:
    doc = Document()
    style_body(doc)
    heading(doc, "Reflexión individual — Semana 1 (L1 + L2)", 0)
    para(
        doc,
        "Entregable: week-1-reflection-<student>.pdf (exportar desde este Word). "
        "Plazo: martes 9 de junio de 2026, 23:59 (Colombia). "
        "Cubre Lección 1 (Access) y Lección 2 (ética y gobernanza). Reemplace student por su nombre.",
    )
    para(doc, "Extensión orientativa del cuerpo: 700–1,000 palabras.")
    doc.add_paragraph()

    section(
        doc,
        "Metadatos",
        "Semana: 1 (L1 + L2) | Estudiante: … | Grupo: G… | Lecturas citadas: … | Conteo de palabras: ~…",
    )
    section(
        doc,
        "R1 Proceso y metodología",
        "¿Qué hizo en los cuadernos l1-introduction y l2-ethics-governance? "
        "¿Qué aprendió sobre Access y sobre auditoría ética?",
    )
    section(
        doc,
        "R2 Justificación técnica",
        "Identifique una decisión técnica (descarga DANE, consulta Overpass u otra) "
        "y explique el por qué.",
    )
    section(
        doc,
        "R3 Enlace con el dominio y los datos",
        "¿Cómo se relacionan GEIH y OSM con su proyecto grupal?",
    )
    section(
        doc,
        "R4 Ética y lecturas",
        "Relacione las lecturas con los aspectos éticos y de gobernanza en los ejercicios prácticos de los cuadernos.",
    )
    section(
        doc,
        "R5 Uso de IA e integridad",
        "¿Qué herramientas de IA utilizó en la semana 1? "
        "¿Cómo influyó en su trabajo y qué aprendió sobre su práctica?",
    )
    section(
        doc,
        "R6 Retroalimentación sobre la semana",
        "¿Qué mejoraría o qué funcionó bien en esta primera semana del curso? "
        "No se califica, únicamente ayuda a ajustar el curso.",
    )
    return doc


def build_reflection_week_2() -> Document:
    doc = Document()
    style_body(doc)
    heading(doc, "Reflexión individual — Semana 2 (L3 + L4)", 0)
    para(
        doc,
        "Entregable: week-2-reflection-<student>.pdf (exportar desde este Word). "
        "Plazo: miércoles 17 de junio de 2026, 23:59 (Colombia). "
        "Cubre Lección 3 (almacenamiento) y Lección 4 (procesamiento). Reemplace student por su nombre.",
    )
    para(doc, "Extensión orientativa del cuerpo (R1–R5): 700–1,000 palabras. R6: 100–200 palabras (sin nota).")
    doc.add_paragraph()

    section(
        doc,
        "Metadatos",
        "Semana: 2 (L3 + L4) | Estudiante: … | Grupo: G… | Lecturas citadas: … | Conteo de palabras: ~… (R1–R5)",
    )
    section(
        doc,
        "R1 Proceso y metodología",
        "¿Qué hizo en los cuadernos l3-storage y l4-processing (2024) y en week-2-group "
        "(harmonización, benchmark, MapReduce, motores)?",
    )
    section(
        doc,
        "R2 Justificación técnica",
        "Explique su elección de partición (L3) y su elección de motor (L4). "
        "Incluya trade-offs y referencias a tiempos medidos en los cuadernos.",
    )
    section(
        doc,
        "R3 Enlace con el dominio y los datos",
        "¿Cómo se relacionan los agregados GEIH harmonizados con la pregunta de decisión de su proyecto? "
        "¿Existe algún output adicional relevante para su proyecto? Si sí, ¿cómo planea implementarlo?",
    )
    section(
        doc,
        "R4 Ética y lecturas",
        "Relacione Dwork y Zuboff cap. 3 con retención en almacenamiento (L3) y supresión al consultar (L4). "
        "Enlace con §6–§7 del PDF de requerimientos del proyecto.",
    )
    section(
        doc,
        "R5 Uso de IA e integridad",
        "¿Qué herramientas de IA utilizó en la semana 2? "
        "¿Cómo influyó en su trabajo y qué aprendió sobre su práctica?",
    )
    section(
        doc,
        "R6 Retroalimentación sobre la semana",
        "¿Qué mejoraría o qué funcionó bien en esta semana del curso? "
        "No se califica, únicamente ayuda a ajustar el curso.",
    )
    return doc


def build_reflection_week_3() -> Document:
    doc = Document()
    style_body(doc)
    heading(doc, "Reflexión individual — Semana 3 (L5 + L6)", 0)
    para(
        doc,
        "Entregable: week-3-reflection-<student>.pdf (exportar desde este Word). "
        "Plazo: miércoles 24 de junio de 2026, 23:59 (Colombia). "
        "Cubre Lección 5 (ingesta y flujos) y Lección 6 (analítica y visualización). Reemplace student por su nombre.",
    )
    para(doc, "Extensión orientativa del cuerpo (R1–R5): 700–1.000 palabras. R6: 100–200 palabras (sin nota).")
    doc.add_paragraph()

    section(
        doc,
        "Metadatos",
        "Semana: 3 (L5 + L6) | Estudiante: … | Grupo: G… | Lecturas citadas: … | Conteo de palabras: ~… (R1–R5)",
    )
    section(
        doc,
        "R1 Proceso y metodología",
        "¿Qué hizo en los cuadernos l5-ingestion y l6-analytics y en week-3-group (curated, Prefect, modelo, gráficos, tablero)?",
    )
    section(
        doc,
        "R2 Justificación técnica",
        "Explique su flujo Prefect (extract → aggregate → validate) y la regresión lineal "
        "que implementó con train/validación/prueba (R², MAE, coeficientes). "
        "Relacione con el contraste lineal vs red pequeña visto en las diapositivas L6 "
        "(explicabilidad vs caja negra).",
    )
    section(
        doc,
        "R3 Enlace con el dominio y los datos",
        "¿Cómo apoyan los agregados curated GEIH su pregunta de decisión? "
        "¿Qué papel tiene (si aplica) el monitor de noticias frente a la capa oficial?",
    )
    section(
        doc,
        "R4 Ética y lecturas",
        "Relacione Zaharia (escala/unificación) y Jarrahi (DCAI) con bitácora, contrato de esquema y rotulado de capas en el tablero.",
    )
    section(
        doc,
        "R5 Uso de IA e integridad",
        "¿Qué herramientas de IA utilizó en la semana 3? "
        "¿Cómo influyó en su trabajo y qué aprendió sobre su práctica?",
    )
    section(
        doc,
        "R6 Retroalimentación sobre la semana",
        "¿Qué mejoraría o qué funcionó bien en esta semana del curso? "
        "No se califica, únicamente ayuda a ajustar el curso.",
    )
    return doc


def build_project_report() -> Document:
    doc = Document()
    style_body(doc)
    heading(doc, "Informe del proyecto — Electiva Big Data", 0)
    para(doc, "Maestría en Estadística Aplicada · UDENAR · A-2026")
    para(
        doc,
        "Consolide el trabajo de las semanas 1–3 (puede reutilizar texto e imágenes de "
        "project_requirements y data_architecture). Exporte a PDF como "
        "project-report-<group_id>.pdf e inclúyalo en el ZIP project-<group_id>.zip. "
        "Plazo del ZIP: martes 30 de junio de 2026, 23:59 (Colombia). "
        "Extensión orientativa del cuerpo (secciones 1–7): 2.000–3.000 palabras "
        "(sin contar figuras, canvas, diagramas ni tablas). "
        "Mantenga los títulos de sección exactamente como aparecen abajo.",
    )
    doc.add_paragraph()

    sections_data = [
        (
            "1. Metadatos del grupo",
            "Id del grupo, integrantes, arquetipo, pregunta de decisión en una oración, conteo aproximado de palabras (secciones 1–7).",
        ),
        (
            "2. Requerimientos del proyecto y modelo canvas",
            "Interesados, requerimientos, objetivos, conjunto de datos. "
            "Incluya el **canvas** (imagen), describalo y liste los requerimientos del proyecto incluyendo los aspectos éticos y de gobernanza.",
        ),
        (
            "3. Pipeline de big data",
            "Figura y narrativa del recorrido completo (fuentes → almacenamiento → procesamiento → "
            "ingesta → analítica → decisión). Indique como se implementaron las diferentes etapas y como dicha implementación cumple con los requerimientos del proyecto.",
        ),
        (
            "4. Arquitectura de datos",
            "Diagrama y explicación: harmonización, partición del lakehouse, ingesta gobernada y herramientas utilizadas.",
        ),
        (
            "5. Analítica y respuesta a la pregunta de decisión",
            "Modelos, métricas, gráficos/tablero; qué apoyan o no la decisión. Indique como se implementaron los modelos y como dicha implementación cumple con los requerimientos del proyecto.",
        ),
        (
            "6. Limitaciones y recomendación",
            "Límites del dato y del modelo; recomendación concreta para el interesado que decide.",
        ),
        (
            "7. Contribución del grupo y uso de IA",
            "Rol de cada integrante; herramientas de IA usadas en el proyecto y cómo influyeron; "
            "declaración de integridad. Si optan por publicación web, confirmen autorización aquí.",
        ),
    ]
    for title, prompt in sections_data:
        section(doc, title, prompt)
        if title.startswith("2."):
            para(doc, "[Espacio para pegar el canvas — inserte imagen en Word]")
            for _ in range(3):
                doc.add_paragraph()
        if title.startswith("4."):
            para(doc, "[Espacio para pegar el diagrama de arquitectura]")
            for _ in range(3):
                doc.add_paragraph()
    return doc


def build_project_reflection() -> Document:
    doc = Document()
    style_body(doc)
    heading(doc, "Reflexión individual — Proyecto final", 0)
    para(
        doc,
        "Entregable: project-reflection-<student>.pdf (exportar desde este Word). "
        "Plazo: miércoles 1 de julio de 2026, 23:59 (Colombia). "
        "Cubre el proyecto integrador (semanas 1–3 y entrega final). Reemplace student por su nombre.",
    )
    para(doc, "Extensión orientativa del cuerpo (R1–R5): 800–1.200 palabras. R6: 100–200 palabras (sin nota).")
    doc.add_paragraph()

    section(
        doc,
        "Metadatos",
        "Proyecto final | Estudiante: … | Grupo: G… | Lecturas citadas: … | Conteo de palabras: ~… (R1–R5)",
    )
    section(
        doc,
        "R1 Proceso y metodología",
        "¿Qué hizo su grupo en el pipeline completo (lakehouse, curated, Prefect, modelos, tablero) "
        "y cuál fue su aporte personal?",
    )
    section(
        doc,
        "R2 Justificación técnica",
        "Defienda dos decisiones técnicas del proyecto (partición, motor, modelo lineal y explicabilidad, "
        "ingesta, etc.) con trade-offs y evidencia de los cuadernos.",
    )
    section(
        doc,
        "R3 Enlace con el dominio y los datos",
        "¿Cómo los agregados y el análisis responden a la pregunta de decisión? "
        "¿Qué capa de evidencia mostraría a quien decide?",
    )
    section(
        doc,
        "R4 Ética y lecturas",
        "Relacione el marco del curso (Zuboff, Dwork, gobernanza L2–L6) con los mecanismos "
        "que aplicó su grupo (bitácora, contrato, k-anon, capas en el tablero).",
    )
    section(
        doc,
        "R5 Uso de IA e integridad",
        "¿Qué herramientas de IA utilizó en el proyecto? "
        "¿Cómo influyó en su trabajo y qué aprendió sobre su práctica?",
    )
    section(
        doc,
        "R6 Retroalimentación sobre el curso",
        "¿Qué mejoraría o qué funcionó bien en el curso? No se califica.",
    )
    return doc


def build_learning_journal_week_1() -> Document:
    doc = Document()
    style_body(doc)

    heading(doc, "Diario de aprendizaje — Semana 1", 0)
    para(doc, "Maestría en Estadística Aplicada · UDENAR · A-2026")
    para(doc, "Electiva I · Big Data")
    para(doc, "Lecturas previas al sábado 6 de junio de 2026 (L1 + L2)")

    doc.add_paragraph()
    heading(doc, "Instrucciones", 1)
    bullet_list(
        doc,
        [
            "Este cuaderno es personal.",
            "Su función es guiar la lectura de los tres textos antes del sábado. "
            "Lleve sus respuestas a la discusión grupal de lecturas y a la plenaria "
            "(~1 h 15 min).",
            "Escriba en español o en inglés (elija uno y manténgalo en todo el documento).",
            "Extensión orientativa: ~600–900 palabras en total (comprensión + reflexión).",
            "Las lecturas se comparten a través del correo electrónico del curso.",
        ],
    )

    heading(doc, "Lecturas de la semana", 1)
    for title, detail in [
        (
            "boyd & Crawford (2012)",
            "“Critical Questions for Big Data”. Archivo: boyd-crawford-2012.pdf",
        ),
        (
            "Zuboff (2019)",
            "La era del capitalismo de la vigilancia — Capítulo 1. "
            "Archivo: zuboff-es-2019.pdf (edición en español) o "
            "zuboff-en-2019.pdf (edición en inglés).",
        ),
        (
            "Mittelstadt et al. (2016)",
            "“The ethics of algorithms: Mapping the debate”. "
            "Archivo: mittelstadt-et-al-2016.pdf",
        ),
    ]:
        reading_bullet(doc, title, detail)

    heading(doc, "Parte A — Lectura 1: Introducción a Big Data", 1)
    para(
        doc,
        "Lea boyd & Crawford y Zuboff Cap. 1. Las preguntas de comprensión "
        "verifican que captó las ideas del texto; las de reflexión temporal "
        "piden comparar el momento en que se escribió con la situación actual "
        "(2026).",
    )

    heading(doc, "A. boyd & Crawford (2012)", 2)
    question(
        doc,
        "A1",
        "En dos o tres oraciones: ¿cuál es la tesis principal del artículo? "
        "¿Qué critican los autores del discurso dominante sobre “Big Data”?",
    )
    question(
        doc,
        "A2",
        "Enumere tres “preguntas críticas” que propone el artículo y explique "
        "una de ellas con sus propias palabras.",
    )
    question(
        doc,
        "A3",
        "¿Qué distingue el artículo entre “más datos” y “mejor comprensión”? "
        "Cite una idea concreta del texto.",
    )
    question(
        doc,
        "A4",
        "Responda una o más “preguntas críticas” de acuerdo a la evolución "
        "tecnológica desde la publicación del artículo (i.e., 2012).",
    )
    question(
        doc,
        "A5",
        "¿Qué limitación o riesgo del “Big Data” señala el artículo que sigue "
        "siendo relevante hoy?",
    )

    heading(doc, "B. Zuboff — Capítulo 1 (2019)", 2)
    question(
        doc,
        "B1",
        "En un párrafo: ¿qué es el “capitalismo de la vigilancia” según Zuboff?",
    )
    question(
        doc,
        "B2",
        "¿Qué relación describe el capítulo entre datos conductuales, predicción "
        "y poder económico? Resuma con sus palabras.",
    )
    question(
        doc,
        "B3",
        "¿En qué se diferencia, según el texto, mejorar un servicio para el "
        "usuario de extraer valor comercial de los datos generados?",
    )
    question(
        doc,
        "B4",
        "Anote una afirmación del capítulo que le resulte convincente y otra "
        "que le genere dudas. Explique brevemente cada una.",
    )

    heading(doc, "Parte B — Lectura 2: Ética y gobernanza de datos", 1)

    heading(doc, "C. Mittelstadt et al. (2016)", 2)
    question(
        doc,
        "C1",
        "¿Qué problema central intenta ordenar el artículo? ¿Qué tipo de "
        "“mapa” del debate propone?",
    )
    question(
        doc,
        "C2",
        "Nombre dos tipos de daño o de preocupación ética del artículo y "
        "explique uno con sus palabras.",
    )
    question(
        doc,
        "C3",
        "¿Qué actores o partes interesadas menciona el texto (p. ej. "
        "desarrolladores, reguladores, personas afectadas)? ¿Quién falta "
        "en el mapa, si alguien falta?",
    )
    question(
        doc,
        "C4",
        "Elija un término del artículo (responsabilidad, transparencia, "
        "equidad, autonomía, explicabilidad, etc.) y defínalo como lo "
        "entiende usted después de leer.",
    )

    heading(doc, "D. Síntesis entre lecturas", 2)
    question(
        doc,
        "D1",
        "En un párrafo: ¿cómo se complementan boyd & Crawford, Zuboff y "
        "Mittelstadt et al. al pensar en datos a gran escala?",
    )
    question(
        doc,
        "D2",
        "¿En qué coinciden los tres textos sobre poder, contexto o límites "
        "de los datos? ¿En qué se separan?",
    )
    question(
        doc,
        "D3",
        "Escriba dos o tres preguntas para la discusión grupal del sábado "
        "sobre las lecturas.",
    )

    heading(doc, "Notas libres", 1)
    para(doc, "Espacio opcional durante la discusión del sábado:")
    for _ in range(6):
        doc.add_paragraph()

    return doc


def build_learning_journal_week_2() -> Document:
    doc = Document()
    style_body(doc)

    heading(doc, "Diario de aprendizaje — Semana 2", 0)
    para(doc, "Maestría en Estadística Aplicada · UDENAR · A-2026")
    para(doc, "Electiva I · Big Data")
    para(doc, "Lecturas previas al sábado 13 de junio de 2026 (L3 + L4)")

    doc.add_paragraph()
    heading(doc, "Instrucciones", 1)
    bullet_list(
        doc,
        [
            "Este cuaderno es personal.",
            "Guía la lectura de los textos antes del sábado. "
            "Lleve sus notas a la discusión grupal de lecturas y a la plenaria.",
            "Extensión orientativa: ~500–700 palabras en total.",
        ],
    )

    heading(doc, "Lecturas (semana 2)", 1)
    for title, detail in [
        (
            "Zuboff (2019)",
            "La era del capitalismo de la vigilancia — **Capítulo 3** "
            "Mismo archivo zuboff-es-2019.pdf o zuboff-en-2019.pdf de la semana 1.",
        ),
        (
            "Dwork (2006)",
            "“Differential privacy”. Archivo: dwork-2006.pdf. ",
        ),
    ]:
        reading_bullet(doc, title, detail)

    heading(doc, "Parte A — Zuboff, Capítulo 3", 1)
    question(
        doc,
        "A1",
        "¿Qué descubrimiento describe Zuboff en Google (inicios de los 2000)? "
        "¿Qué cambió en la forma de valorar los datos de búsqueda y navegación?",
    )
    question(
        doc,
        "A2",
        "¿En qué se diferencia, según el capítulo, mejorar un servicio para el "
        "usuario de convertir datos conductuales en predicción vendible?",
    )
    question(
        doc,
        "A3",
        "Compare en un párrafo el régimen de datos de Google (cap. 3) con el de "
        "microdatos públicos GEIH/DANE que usa el curso. ¿Quién se beneficia en "
        "cada caso?",
    )
    question(
        doc,
        "A4",
        "Anote una afirmación del capítulo que le resulte convincente y otra que "
        "le genere dudas. Explique brevemente.",
    )

    heading(doc, "Parte B — Dwork (2006), privacidad diferencial", 1)
    para(
        doc,
        "¿Qué promete la privacidad diferencial y "
        "por qué importa al procesar microdatos?",
    )
    question(
        doc,
        "B1",
        "¿Qué problema intenta resolver la privacidad diferencial que no resuelve "
        "solo “quitar nombres” o publicar agregados?",
    )
    question(
        doc,
        "B2",
        "¿Qué significa el trade-off entre privacidad y utilidad "
        "analítica? ¿Por qué no se puede tener “privacidad perfecta” y “datos "
        "idénticos sin ruido” a la vez?",
    )
    question(
        doc,
        "B3",
        "Relacione Dwork con la semana 1: cuasi-identificadores en GEIH y riesgo "
        "de divulgación en tablas agregadas.",
    )
    question(
        doc,
        "B4",
        "Si su grupo publicara un informe departamental (conteos, mapas, rankings), "
        "¿qué precaución de privacidad propondría inspirada en esta lectura?",
    )

    heading(doc, "Notas libres", 1)
    para(doc, "Espacio para citas, dudas o conexiones con su proyecto grupal:")
    for _ in range(6):
        doc.add_paragraph()

    return doc


def build_learning_journal_week_3() -> Document:
    doc = Document()
    style_body(doc)

    heading(doc, "Diario de aprendizaje — Semana 3", 0)
    para(doc, "Maestría en Estadística Aplicada · UDENAR · A-2026")
    para(doc, "Electiva I · Big Data")
    para(doc, "Lecturas previas al sábado 20 de junio de 2026 (L5 + L6)")

    doc.add_paragraph()
    heading(doc, "Instrucciones", 1)
    bullet_list(
        doc,
        [
            "Este cuaderno es personal.",
            "Guía la lectura de los textos antes del sábado. "
            "Lleve sus notas a la discusión grupal de lecturas y a la plenaria.",
            "Extensión orientativa: ~500–700 palabras en total.",
            "Las lecturas se comparten a través del correo electrónico del curso.",
        ],
    )

    heading(doc, "Lecturas (semana 3)", 1)
    for title, detail in [
        (
            "Zaharia et al. (2016)",
            "“Apache Spark: a unified engine for big data processing”. "
            "Archivo: zaharia-2016.pdf (adjunto al correo).",
        ),
        (
            "Jarrahi et al. (2023)",
            "“The Principles of Data-Centric AI”. "
            "Archivo: jarrahi-2023.pdf (adjunto al correo).",
        ),
    ]:
        reading_bullet(doc, title, detail)

    heading(doc, "Parte A — Zaharia et al. (2016), Apache Spark", 1)
    para(
        doc,
        "Esta lectura apoya las lecciones 5 (ingestión y flujo de trabajo) y 6 "
        "(analítica avanzada y visualización): procesamiento distribuido y ML sobre "
        "datos grandes.",
    )
    question(
        doc,
        "A1",
        "En dos o tres oraciones: ¿qué problema intenta resolver Spark frente a "
        "herramientas separadas para batch, consultas interactivas, streaming y ML?",
    )
    question(
        doc,
        "A2",
        "¿Qué ventajas describe el artículo al unificar estos workloads en un solo "
        "motor?",
    )
    question(
        doc,
        "A3",
        "¿Qué limitación o coste de operar a escala (infraestructura, complejidad, "
        "latencia, debugging) menciona o implica el artículo?",
    )

    heading(doc, "Parte B — Jarrahi et al. (2023), IA centrada en los datos (DCAI)", 1)
    para(
        doc,
        "Hacia dónde va la práctica con datos e IA, más allá del hype de “solo mejorar el modelo”.",
    )
    question(
        doc,
        "B1",
        "¿Qué critica el artículo del enfoque centrado en el modelo (model-centric AI)? "
        "¿Qué propone en su lugar el enfoque centrado en los datos (DCAI)? ",
    )
    question(
        doc,
        "B2",
        "El artículo formula seis principios de DCAI. Elija dos y explíquelos con un "
        "ejemplo concreto de su proyecto.",
    )
    question(
        doc,
        "B3",
        "¿Qué significa que los datos son un constructo sociotécnico y "
        "que el trabajo con datos es centrado en las personas? ",
    )

    heading(doc, "Notas libres", 1)
    para(doc, "Espacio para citas, dudas o conexiones con su proyecto grupal:")
    for _ in range(6):
        doc.add_paragraph()

    return doc


# --- write targets -----------------------------------------------------------


def write_project_requirements() -> tuple[Path, Path]:
    return save_mirror(
        build_project_requirements(),
        PROJECT_REQUIREMENTS_ASSETS,
        PROJECT_REQUIREMENTS_WEEK1,
    )


def write_data_architecture() -> tuple[Path, Path]:
    return save_mirror(
        build_data_architecture(),
        DATA_ARCHITECTURE_ASSETS,
        DATA_ARCHITECTURE_WEEK2,
    )


def write_reflection_week_1() -> tuple[Path, Path]:
    return save_mirror(
        build_reflection_week_1(),
        REFLECTION_WEEK_1_ASSETS,
        REFLECTION_WEEK_1_WEEK1,
    )


def write_reflection_week_2() -> tuple[Path, Path]:
    return save_mirror(
        build_reflection_week_2(),
        REFLECTION_WEEK_2_ASSETS,
        REFLECTION_WEEK_2_WEEK2,
    )


def write_reflection_week_3() -> tuple[Path, Path]:
    return save_mirror(
        build_reflection_week_3(),
        REFLECTION_WEEK_3_ASSETS,
        REFLECTION_WEEK_3_WEEK3,
    )


def write_project_report() -> tuple[Path, Path]:
    return save_mirror(
        build_project_report(),
        PROJECT_REPORT_ASSETS,
        PROJECT_REPORT_PROJECT,
    )


def write_project_reflection() -> tuple[Path, Path]:
    return save_mirror(
        build_project_reflection(),
        PROJECT_REFLECTION_ASSETS,
        PROJECT_REFLECTION_PROJECT,
    )


def write_learning_journal_week_1() -> tuple[Path, Path]:
    return save_mirror(
        build_learning_journal_week_1(),
        LEARNING_JOURNAL_WEEK_1_ASSETS,
        LEARNING_JOURNAL_WEEK_1,
    )


def write_learning_journal_week_2() -> tuple[Path, Path]:
    return save_mirror(
        build_learning_journal_week_2(),
        LEARNING_JOURNAL_WEEK_2_ASSETS,
        LEARNING_JOURNAL_WEEK_2,
    )


def write_learning_journal_week_3() -> tuple[Path, Path]:
    return save_mirror(
        build_learning_journal_week_3(),
        LEARNING_JOURNAL_WEEK_3_ASSETS,
        LEARNING_JOURNAL_WEEK_3,
    )


DOCUMENT_WRITERS: dict[str, Callable[[], list[Path]]] = {
    "project_requirements": lambda: list(write_project_requirements()),
    "data_architecture": lambda: list(write_data_architecture()),
    "reflection_week_1": lambda: list(write_reflection_week_1()),
    "reflection_week_2": lambda: list(write_reflection_week_2()),
    "reflection_week_3": lambda: list(write_reflection_week_3()),
    "project_report": lambda: list(write_project_report()),
    "project_reflection": lambda: list(write_project_reflection()),
    "learning_journal_week_1": lambda: list(write_learning_journal_week_1()),
    "learning_journal_week_2": lambda: list(write_learning_journal_week_2()),
    "learning_journal_week_3": lambda: list(write_learning_journal_week_3()),
}


def build_all() -> list[Path]:
    paths: list[Path] = []
    for writer in DOCUMENT_WRITERS.values():
        paths.extend(writer())
    return paths


def main() -> None:
    parser = argparse.ArgumentParser(description="Build Big Data course Word documents.")
    parser.add_argument(
        "document",
        nargs="?",
        choices=tuple(DOCUMENT_WRITERS.keys()),
        help="Document to build (omit with --all)",
    )
    parser.add_argument(
        "--all",
        action="store_true",
        help="Build every course Word document",
    )
    args = parser.parse_args()

    if args.all:
        paths = build_all()
    elif args.document:
        paths = DOCUMENT_WRITERS[args.document]()
    else:
        parser.error("specify a document name or --all")

    for path in paths:
        print(f"Wrote {path}")


if __name__ == "__main__":
    main()
