"""Shared helpers for Big Data course Word documents (python-docx)."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Pt

REPO = Path(__file__).resolve().parents[2]
COURSE = "26-udenar-big-data"
ASSETS_DIR = REPO / "assets" / "documents" / COURSE
COURSE_WS = REPO / "work-space/teaching/big-data/course"
WEEK1_DIR = COURSE_WS / "week-1"
WEEK2_DIR = COURSE_WS / "week-2"


def style_body(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.15
    pf.space_after = Pt(6)


def heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def para(doc: Document, text: str, *, bold: bool = False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold


def section(doc: Document, title: str, prompt: str) -> None:
    heading(doc, title, 2)
    para(doc, prompt)
    doc.add_paragraph()


def question(doc: Document, number: str, text: str) -> None:
    p = doc.add_paragraph()
    p.add_run(f"{number}. ").bold = True
    p.add_run(text)
    doc.add_paragraph("Respuesta:", style="Intense Quote")
    for _ in range(3):
        doc.add_paragraph()


def bullet_list(doc: Document, lines: list[str]) -> None:
    for line in lines:
        doc.add_paragraph(line, style="List Bullet")


def reading_bullet(doc: Document, title: str, detail: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(f"{title}. ").bold = True
    p.add_run(detail)


def save_doc(doc: Document, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path


def save_mirror(doc: Document, assets_path: Path, mirror_path: Path) -> tuple[Path, Path]:
    save_doc(doc, assets_path)
    save_doc(doc, mirror_path)
    return assets_path, mirror_path
