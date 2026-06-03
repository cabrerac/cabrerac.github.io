#!/usr/bin/env python3
"""Fill Udenar Programa Asignatura docx from planning content (Spanish)."""

from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document

REPO = Path(__file__).resolve().parents[2]
TEMPLATE = (
    REPO
    / "work-space/teaching/big-data/administration"
    / "Formato de Programa Asignatura ELECTIVA I -BIG DATA - MESA A-2026.docx"
)
OUTPUT = TEMPLATE.with_name(
    "Formato de Programa Asignatura ELECTIVA I -BIG DATA - MESA A-2026-FILLED.docx"
)

# --- Edit these if needed ---
DOCENTE = "Christian Cabrera Jojoa"
DOCENTE_ID = ""  # cédula — completar en Word si vacío
DOCENTE_EMAIL = "chc79@cam.ac.uk"
DOCENTE_MOVIL = ""  # WhatsApp — completar en Word si vacío
CODIGO_ASIGNATURA = ""
FECHA_ACTUALIZACION = "02-06-2026"
REVISION_POR = "Christian Cabrera Jojoa"
DIRECTOR_PROGRAMA = ""

JUSTIFICACION = """Los organismos públicos, centros de investigación y empresas generan y usan volúmenes crecientes de datos con exigencias de velocidad, variedad y trazabilidad que superan el análisis estadístico tradicional en una sola máquina o con una sola herramienta. La Maestría en Estadística Aplicada forma profesionales que deben enmarcar problemas, elegir arquitecturas y formatos, procesar datos a escala moderada con herramientas actuales y comunicar hallazgos sin perder de vista límites, sesgos y gobernanza.

Esta electiva introduce el proceso de big data (acceso, evaluación, tratamiento, almacenamiento, procesamiento, ingesta y analítica) sobre datos reales de contexto colombiano (microdatos GEIH del DANE), con énfasis en decisiones técnicas justificadas, ética desde el inicio del ciclo y proyecto integrador en equipo. El enfoque evita el tecnocentrismo: las herramientas sirven a un problema y a unos decisores identificados, en línea con la formación aplicada de la maestría."""

OBJETIVO_GENERAL = """Desarrollar competencias para identificar situaciones de big data, aplicar una metodología reproducible de extremo a extremo, implementar soluciones con herramientas modernas de código abierto (notebooks en Colab y entorno local) y analizar implicaciones éticas y de gobernanza de las decisiones de diseño, mediante actividades formativas por módulo y un proyecto integrador grupal con reflexión individual."""

OBJETIVOS_ESPECIFICOS = """Al finalizar, el estudiante estará en capacidad de:

• Caracterizar un problema con las dimensiones de big data (volumen, velocidad, variedad, veracidad, etc.) y criterios de éxito explícitos.
• Aplicar la metodología del curso (tres A: acceso, evaluación, respuesta) a lo largo de un pipeline documentado.
• Comparar patrones de arquitectura (nube, borde, lago/lakehouse, batch vs interactivo) y justificar compensaciones para un caso dado.
• Fundamentar decisiones de almacenamiento y formatos (p. ej. Parquet, particionamiento, motores pandas / Polars / DuckDB / Spark local).
• Aplicar técnicas de procesamiento y control de calidad acordes a la escala del laboratorio e interpretar resultados en contexto.
• Diseñar o criticar flujos de ingesta y orquestación (ETL/ELT; batch vs streaming conceptual).
• Producir analítica y visualizaciones orientadas a una pregunta de decisión, con narrativa explícita.
• Analizar privacidad, gobernanza y ética vinculadas a las decisiones técnicas de gestión y procesamiento de datos.
• Integrar los módulos 1–6 en un proyecto summativo grupal con aportes individuales documentados."""

METODOLOGIA = """Modalidad híbrida: ~2 h asíncronas por módulo (video de teoría + demostración de laboratorio) + bloque presencial los sábados (07:00–13:00, hora Colombia), tres sábados de enseñanza con dos módulos por jornada, más jornada final de presentaciones (27 de junio de 2026).

Antes de cada sábado: lecturas guiadas (diario personal), videos asíncronos y cuaderno de práctica individual en Google Colab. En el sábado: discusión de lecturas (~1 h 15 min), comparación de soluciones, clínica de proyecto y consolidación en Colab grupal (rotación de escriba). Después del sábado: tarea grupal con auto-verificaciones; reflexión individual en Markdown (Moodle); plazo miércoles 23:59 (hora Colombia).

Equipos estables (~5 grupos) sobre escenario común (analítica para decisiones de política pública con GEIH) y documento de requerimientos del proyecto (PDF). Recursos: sitio web del curso, Moodle, Microsoft Teams. Stack: Python (pandas, Polars, DuckDB, PySpark local, Prefect, dbt-core, scikit-learn); sin clúster obligatorio."""

CRITERIOS = """Seis evaluaciones formativas (10 % cada una) + proyecto final integrador (40 %).

Cada módulo formativo (10 %): 60 % cuaderno grupal (artefacto técnico con verificaciones automáticas) + 40 % reflexión individual (proceso, justificación técnica, dominio, ética y lecturas, uso de IA).

Proyecto final (40 %): 70 % entregables grupales (cuaderno integrado, diagrama, informe de decisión, presentación) + 30 % reflexión individual integradora.

Discusión de lecturas: participación esperada; retroalimentación sin nota. Integridad académica: declaración de uso de IA y registro de contribuciones por integrante."""

COMPETENCIAS = """1. Enmarcado analítico — Problema, actores, restricciones y criterios de éxito.
2. Ejecución técnica — Pipeline coherente dentro de los recursos del curso.
3. Comunicación — Explicación de diseño y hallazgos (visual, escrita, oral).
4. Práctica responsable — Ética y gobernanza en gestión y procesamiento de datos."""

REFERENCIAS = """LECTURAS ASIGNADAS POR SEMANA

Semana 1 (Módulos 1–2): boyd & Crawford (2012); Zuboff (2019) Cap. 1; Mittelstadt et al. (2016).

Semana 2 (Módulos 3–4): Dwork (2006), privacidad diferencial; Armbrust et al. (2021), lakehouse (lectura orientativa).

Semana 3 (Módulos 5–6): Peng (2011), reproducibilidad; Barocas & Selbst (2016), impacto dispar.

BIBLIOGRAFÍA COMPLEMENTARIA (programa)

Nissenbaum (2004), integridad contextual de la privacidad; Donoho (2017), 50 years of data science; Hand (2020), retos estadísticos del big data; Dean & Ghemawat (2004), MapReduce; Stonebraker et al. (2010), MapReduce vs DBMS; O'Neil (2016), Weapons of math destruction (opcional).

DATOS Y TÉCNICOS

DANE — GEIH (microdatos); documentación Parquet, Polars, DuckDB, Spark, Prefect, dbt. Materiales del curso (web + Moodle). Detalle: planning/reading-list.md."""

MODULOS = [
    (
        "5",
        "Módulo 1. Introducción a big data, metodología y ecosistemas. Dimensiones V; metodología (tres A); arquitecturas; acceso GEIH; requerimientos del proyecto (PDF).",
        "Formativa 10 % (cuaderno grupal + reflexión individual)",
    ),
    (
        "5",
        "Módulo 2. Ética, privacidad y gobernanza. Marco ético; lecturas; auditoría ética; actualización de requerimientos (práctica responsable).",
        "Formativa 10 %",
    ),
    (
        "5",
        "Módulo 3. Almacenamiento y gestión. Parquet; lakehouse; Polars/DuckDB; particionamiento; retención y linaje.",
        "Formativa 10 %",
    ),
    (
        "5",
        "Módulo 4. Procesamiento y análisis. Motores de procesamiento; Spark local; calidad; k-anonimato / privacidad diferencial.",
        "Formativa 10 %",
    ),
    (
        "5",
        "Módulo 5. Ingesta y flujos. ETL/ELT; Prefect; dbt-core; streaming; auditoría y proveniencia.",
        "Formativa 10 %",
    ),
    (
        "5",
        "Módulo 6. Analítica y visualización. Modelos y gráficos; tableros y gobernanza; cierre del hilo ético.",
        "Formativa 10 %",
    ),
    (
        "—",
        "Semana de proyecto e integración (async + trabajo grupal; presentaciones 27 jun 2026).",
        "Summativa 40 %",
    ),
]


def _set_cell_text(cell, text: str) -> None:
    cell.text = text


def fill() -> Path:
    doc = Document(TEMPLATE)

    # Header lines (if present after tables in some templates — here before tables)
    for para in doc.paragraphs:
        if para.text.startswith("Docente:"):
            para.text = f"Docente: {DOCENTE}"
        elif para.text.startswith("Correo"):
            para.text = f"Correo Electrónico: {DOCENTE_EMAIL}"
        elif para.text.startswith("Móvil") or para.text.startswith("M\u00f3vil"):
            para.text = f"Móvil (WhatsApp): {DOCENTE_MOVIL}"

    t0 = doc.tables[0]
    _set_cell_text(t0.rows[3].cells[0], DOCENTE)
    if DOCENTE_ID:
        _set_cell_text(t0.rows[3].cells[3], DOCENTE_ID)
    _set_cell_text(t0.rows[4].cells[1], DOCENTE_EMAIL)
    if CODIGO_ASIGNATURA:
        _set_cell_text(t0.rows[6].cells[1], CODIGO_ASIGNATURA)
    _set_cell_text(t0.rows[9].cells[1], "5")
    _set_cell_text(t0.rows[9].cells[2], "5")
    _set_cell_text(t0.rows[9].cells[3], "20")
    _set_cell_text(t0.rows[9].cells[4], "30")
    _set_cell_text(t0.rows[11].cells[0], f"Fecha última Actualización del Programa Temático: {FECHA_ACTUALIZACION}")
    _set_cell_text(t0.rows[11].cells[1], REVISION_POR)
    if DIRECTOR_PROGRAMA:
        _set_cell_text(t0.rows[11].cells[3], DIRECTOR_PROGRAMA)

    _set_cell_text(doc.tables[1].rows[1].cells[0], JUSTIFICACION)
    _set_cell_text(doc.tables[2].rows[2].cells[0], OBJETIVO_GENERAL)
    _set_cell_text(doc.tables[2].rows[4].cells[0], OBJETIVOS_ESPECIFICOS)
    _set_cell_text(doc.tables[3].rows[1].cells[0], METODOLOGIA)
    _set_cell_text(doc.tables[4].rows[1].cells[0], CRITERIOS)
    _set_cell_text(doc.tables[6].rows[1].cells[0], COMPETENCIAS)
    _set_cell_text(doc.tables[7].rows[1].cells[0], REFERENCIAS)

    t5 = doc.tables[5]
    # Clear placeholder row 2 content then add module rows
    _set_cell_text(t5.rows[2].cells[0], "")
    _set_cell_text(t5.rows[2].cells[1], "")
    _set_cell_text(t5.rows[2].cells[2], "")

    for horas, tema, eval_form in MODULOS:
        row = t5.add_row()
        _set_cell_text(row.cells[0], horas)
        _set_cell_text(row.cells[1], tema)
        _set_cell_text(row.cells[2], eval_form)

    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    path = fill()
    print(f"Wrote {path}")
