#!/usr/bin/env python3
"""Build the Spanish project outline (definition + three archetypes) for the Big Data elective.

Outputs (canonical for the site):
  assets/documents/26-udenar-big-data/project-definition-big-data.pdf
  assets/documents/26-udenar-big-data/project-definition-big-data.docx

Mirror:
  work-space/teaching/big-data/course/project-definition-big-data.pdf
  work-space/teaching/big-data/course/project-definition-big-data.docx

Regenerate:
  python scripts/big-data-course/build_project_outline.py
"""

from __future__ import annotations

import html
from pathlib import Path

import pymupdf as fitz
from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Pt

REPO = Path(__file__).resolve().parents[2]
COURSE = "26-udenar-big-data"
ASSETS_DIR = REPO / "assets" / "documents" / COURSE
COURSE_DIR = REPO / "work-space/teaching/big-data/course"

PDF_ASSETS = ASSETS_DIR / "project-definition-big-data.pdf"
DOCX_ASSETS = ASSETS_DIR / "project-definition-big-data.docx"
PDF_COURSE = COURSE_DIR / "project-definition-big-data.pdf"
DOCX_COURSE = COURSE_DIR / "project-definition-big-data.docx"

TITLE = "Proyecto integrador — Electiva Big Data"
SUBTITLE = "Maestría en Estadística Aplicada · Universidad de Nariño · A-2026"

CSS = """
body { font-family: Helvetica, Arial, sans-serif; font-size: 11pt; line-height: 1.35; }
h1 { font-size: 18pt; margin-bottom: 4pt; }
h2 { font-size: 13pt; margin-top: 14pt; margin-bottom: 6pt; }
h3 { font-size: 11.5pt; margin-top: 10pt; margin-bottom: 4pt; }
p { margin: 0 0 8pt 0; }
ul { margin: 0 0 8pt 0; padding-left: 18pt; }
li { margin-bottom: 4pt; }
.subtitle { color: #333; margin-bottom: 12pt; }
.box { background: #f5f5f5; padding: 8pt; margin: 8pt 0; }
"""

SECTIONS: list[tuple[str, list[str]]] = [
    (
        "1. De qué trata el proyecto",
        [
            "Ustedes y su grupo forman un equipo de analítica que apoya decisiones de "
            "política pública en Colombia. El trabajo del curso es construir un pipeline "
            "de big data reproducible: Access, Assess y Address sobre datos reales.",
            "La pregunta concreta, los interesados y el subconjunto de datos los define "
            "cada grupo. Este documento fija el marco común y las tres opciones de enfoque "
            "para el proyecto. "
        ],
    ),
    (
        "2. Datos disponibles",
        [
            "Dataset obligatorio: microdatos GEIH del DANE (Gran Encuesta Integrada de "
            "Hogares). En la Lección 1 acceden a archivos mensuales; a lo largo del "
            "curso trabajan con particiones harmonizadas incluyendo rangos de años mayores (p. ej. 2022–2025).",
            "El dataset GEIH trae información de mercado laboral a nivel persona-mes: empleo, "
            "informalidad, ocupación, ingresos (según año), geografía (departamento, "
            "área), pesos de expansión y otras variables demográficas relevantes para el proyecto.",
            "Las preguntas de decisión deben responderse con agregados del dataset GEIH "
            "(departamento, segmento, periodo). No se orienta el proyecto a identificar personas ni a usar datos que la encuesta no contiene.",
        ],
    ),
    (
        "3. Fuentes adicionales (opcional, recomendado)",
        [
            "Animamos a enriquecer el pipeline con otras fuentes abiertas cuando aporten "
            "a la pregunta de decisión. Ejemplos previstos en el curso: OpenStreetMap (contexto geográfico: servicios, vías, POI agregados por zona), "
            "Tablas administrativas abiertas del DANE u otras entidades públicas. "
            "El dataset GEIH sigue siendo la base común del curso. Las fuentes extra complementan, "
            "no la reemplazan.",
        ],
    ),
    (
        "4. Tres opciones de enfoque (elija uno por grupo)",
        [
            "Todos los grupos usan el dataset GEIH. Lo que cambia es el tipo de decisión "
            "que apoyan. Escoja una de las tres opciones de enfoque.",
        ],
    ),
]

ARCHETYPES: list[tuple[str, str, list[str]]] = [
    (
        "A. Asignación de recursos",
        "Distribuir programas, servicios o inversión entre regiones o segmentos para mejorar el mercado laboral.",
        [
            "Ejemplo: ¿Qué departamentos o segmentos de población deberían ser prioridad "
            "para un programa de empleo este año?",
            "Indicadores típicos: tasas de informalidad, desempleo o inactividad por "
            "departamento, con pesos de expansión.",
            "OpenStreetMap (opcional): agregar por departamento o municipio la densidad de "
            "servicios (salud, educación, transporte) para contrastar prioridades GEIH con "
            "contexto geográfico.",
            "Límite: el dataset GEIH informa prioridades, no registra cupos ni presupuestos de "
            "programas.",
        ],
    ),
    (
        "B. Riesgo / alerta temprana",
        "Identificar zonas o condiciones donde el riesgo laboral aumenta para tomar medidas preventivas.",
        [
            "Ejemplo: ¿Dónde aumentó más la informalidad o el desempleo entre un periodo definido?",
            "Indicadores típicos: cambios en el tiempo por departamento o segmento.",
            "OpenStreetMap (opcional): señalar zonas donde empeoran indicadores GEIH y "
            "faltan vías o POI de apoyo (servicios, conectividad) según mapas agregados.",
            "Límite: el dataset GEIH es encuesta mensual con rezago, no alerta en tiempo real. "
            "Formule la pregunta como tendencia o deterioro relativo, no predicción "
            "instantánea.",
        ],
    ),
    (
        "C. Monitoreo / seguimiento",
        "Seguir un indicador en el tiempo y reportarlo a decisores.",
        [
            "Ejemplo: ¿Cómo evolucionan las brechas regionales en participación laboral "
            "entre un periodo definido?",
            "Indicadores típicos: series anuales o mensuales comparables, con "
            "armonización de códigos entre años.",
            "OpenStreetMap (opcional): fijar una capa geográfica de referencia (p. ej. POI o "
            "red vial por zona) para interpretar tendencias GEIH en el mismo territorio.",
        ],
    ),
]

FOOTER_SECTIONS: list[tuple[str, list[str]]] = [
    (
        "5. Pregunta de referencia del curso",
        [
            "Independiente del enfoque, en los entregables finales cada grupo debe "
            "responder al menos una vez:",
            "«¿Qué dice nuestro pipeline sobre patrones del mercado laboral para un "
            "segmento de población definido, y qué recomendaríamos o no recomendaríamos "
            "con base en ello?»",
        ],
    ),
    (
        "6. Próximos pasos",
        [
            "1. Lea este documento con su grupo y discutanlo durante las sesiones grupales.",
            "2. Elija un enfoque y redacten project_requirements.pdf con la plantilla "
            "Word del curso. Este documento se discutirá en la segunda semana del curso.",
            "3. Si planean usar un dataset adicional, documentenlo en project_requirements.pdf.",
            "4. Alimente la definición y desarrollo del proyecto de acuerdo a lo explorado en las lecciones del curso.",
        ],
    )
]

def _style_body(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.15
    pf.space_after = Pt(6)


def _build_docx() -> Document:
    doc = Document()
    _style_body(doc)
    doc.add_heading(TITLE, 0)
    doc.add_paragraph(SUBTITLE)
    doc.add_paragraph()

    for heading, paragraphs in SECTIONS:
        doc.add_heading(heading, level=1)
        for text in paragraphs:
            doc.add_paragraph(text)

    for title, decision, bullets in ARCHETYPES:
        doc.add_heading(title, level=2)
        doc.add_paragraph(decision)
        for item in bullets:
            doc.add_paragraph(item, style="List Bullet")

    for heading, paragraphs in FOOTER_SECTIONS:
        doc.add_heading(heading, level=1)
        for text in paragraphs:
            if text.startswith("«"):
                doc.add_paragraph(text)
            else:
                doc.add_paragraph(text)

    return doc


def _p(text: str) -> str:
    return f"<p>{html.escape(text)}</p>"


def _li(text: str) -> str:
    return f"<li>{html.escape(text)}</li>"


def _build_html() -> str:
    parts = [
        f"<h1>{html.escape(TITLE)}</h1>",
        f'<p class="subtitle">{html.escape(SUBTITLE)}</p>',
    ]

    for heading, paragraphs in SECTIONS:
        parts.append(f"<h2>{html.escape(heading)}</h2>")
        for text in paragraphs:
            parts.append(_p(text))

    for title, decision, bullets in ARCHETYPES:
        parts.append(f"<h3>{html.escape(title)}</h3>")
        parts.append(_p(decision))
        parts.append("<ul>")
        parts.extend(_li(b) for b in bullets)
        parts.append("</ul>")

    for heading, paragraphs in FOOTER_SECTIONS:
        parts.append(f"<h2>{html.escape(heading)}</h2>")
        for text in paragraphs:
            parts.append(_p(text))

    return f"<body>{''.join(parts)}</body>"


def _write_pdf(path: Path, body_html: str) -> None:
    story = fitz.Story(html=body_html, user_css=CSS)
    mediabox = fitz.paper_rect("a4")
    where = mediabox + (50, 50, -50, -50)
    writer = fitz.DocumentWriter(str(path))
    try:
        more = 1
        while more:
            device = writer.begin_page(mediabox)
            more, _filled = story.place(where)
            story.draw(device, None)
            writer.end_page()
    finally:
        writer.close()


def _save_pair(doc: Document, assets_docx: Path, course_docx: Path) -> None:
    assets_docx.parent.mkdir(parents=True, exist_ok=True)
    course_docx.parent.mkdir(parents=True, exist_ok=True)
    doc.save(assets_docx)
    doc.save(course_docx)


def build() -> tuple[Path, Path]:
    doc = _build_docx()
    _save_pair(doc, DOCX_ASSETS, DOCX_COURSE)
    _write_pdf(PDF_ASSETS, _build_html())
    PDF_COURSE.write_bytes(PDF_ASSETS.read_bytes())
    return PDF_ASSETS, DOCX_ASSETS


if __name__ == "__main__":
    pdf_path, docx_path = build()
    print(f"Wrote {pdf_path}")
    print(f"Wrote {docx_path}")
    print(f"Mirror {PDF_COURSE}")
    print(f"Mirror {DOCX_COURSE}")
