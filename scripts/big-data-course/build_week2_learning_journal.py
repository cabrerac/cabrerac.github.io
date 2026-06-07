#!/usr/bin/env python3
"""Scaffold for Week 2 learning journal (Spanish).

Canonical student file: course/week-2/learning-journal-week-2.docx
Regenerating overwrites that path — merge manual edits back into this script if needed.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Pt

REPO_ROOT = Path(__file__).resolve().parents[2]
OUT_PATH = (
    REPO_ROOT
    / "work-space/teaching/big-data/course/week-2/learning-journal-week-2.docx"
)


def _style_body(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.15
    pf.space_after = Pt(6)


def _heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def _para(doc: Document, text: str) -> None:
    doc.add_paragraph(text)


def _question(doc: Document, number: str, text: str) -> None:
    p = doc.add_paragraph()
    p.add_run(f"{number}. ").bold = True
    p.add_run(text)
    doc.add_paragraph("Respuesta:", style="Intense Quote")
    for _ in range(3):
        doc.add_paragraph()


def build() -> Path:
    doc = Document()
    _style_body(doc)

    _heading(doc, "Diario de aprendizaje — Semana 2", 0)
    _para(doc, "Maestría en Estadística Aplicada · UDENAR · A-2026")
    _para(doc, "Electiva I · Big Data")
    _para(doc, "Lecturas previas al sábado 13 de junio de 2026 (L3 + L4)")

    doc.add_paragraph()
    _heading(doc, "Instrucciones", 1)
    for line in [
        "Este cuaderno es personal.",
        "Guía la lectura de los textos antes del sábado. "
        "Lleve sus notas a la discusión grupal de lecturas y a la plenaria.",
        "Extensión orientativa: ~500–700 palabras en total.",
    ]:
        doc.add_paragraph(line, style="List Bullet")

    _heading(doc, "Lecturas (semana 2)", 1)
    mandatory = [
        (
            "Zuboff (2019)",
            "La era del capitalismo de la vigilancia — **Capítulo 3** "
            "Mismo archivo zuboff-es-2019.pdf o zuboff-en-2019.pdf de la semana 1.",
        ),
        (
            "Dwork (2006)",
            "“Differential privacy”. Archivo: dwork-2006.pdf. ",
        ),
    ]
    for title, detail in mandatory:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(f"{title}. ").bold = True
        p.add_run(detail)

    _heading(doc, "Parte A — Zuboff, Capítulo 3", 1)
    _question(
        doc,
        "A1",
        "¿Qué descubrimiento describe Zuboff en Google (inicios de los 2000)? "
        "¿Qué cambió en la forma de valorar los datos de búsqueda y navegación?",
    )
    _question(
        doc,
        "A2",
        "¿En qué se diferencia, según el capítulo, mejorar un servicio para el "
        "usuario de convertir datos conductuales en predicción vendible?",
    )
    _question(
        doc,
        "A3",
        "Compare en un párrafo el régimen de datos de Google (cap. 3) con el de "
        "microdatos públicos GEIH/DANE que usa el curso. ¿Quién se beneficia en "
        "cada caso?",
    )
    _question(
        doc,
        "A4",
        "Anote una afirmación del capítulo que le resulte convincente y otra que "
        "le genere dudas. Explique brevemente.",
    )

    _heading(doc, "Parte B — Dwork (2006), privacidad diferencial", 1)
    _para(
        doc,
        "¿Qué promete la privacidad diferencial y "
        "por qué importa al procesar microdatos?",
    )
    _question(
        doc,
        "B1",
        "¿Qué problema intenta resolver la privacidad diferencial que no resuelve "
        "solo “quitar nombres” o publicar agregados?",
    )
    _question(
        doc,
        "B2",
        "¿Qué significa el trade-off entre privacidad y utilidad "
        "analítica? ¿Por qué no se puede tener “privacidad perfecta” y “datos "
        "idénticos sin ruido” a la vez?",
    )
    _question(
        doc,
        "B3",
        "Relacione Dwork con la semana 1: cuasi-identificadores en GEIH y riesgo "
        "de divulgación en tablas agregadas.",
    )
    _question(
        doc,
        "B4",
        "Si su grupo publicara un informe departamental (conteos, mapas, rankings), "
        "¿qué precaución de privacidad propondría inspirada en esta lectura?",
    )

    _heading(doc, "Notas libres", 1)
    _para(doc, "Espacio para citas, dudas o conexiones con su proyecto grupal:")
    for _ in range(6):
        doc.add_paragraph()

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_PATH)
    return OUT_PATH


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path}")
